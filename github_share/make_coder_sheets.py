#!/usr/bin/env python3
# Builds the Excel coding workbooks and text files for each coding round.

from __future__ import annotations

import argparse
import difflib
import itertools
import json
import random
import re
from pathlib import Path

import pandas as pd

DIMS = ["T", "R", "A", "C", "E"]
LEVELS = {
    "T": ["T1_Not_Required_or_Not_Mentioned", "T2_Encouraged_or_Voluntary",
          "T3_Basic_Disclosure_Required", "T4_Detailed_Traceability_Required"],
    "R": ["R1_Not_Mentioned", "R2_Human_Ownership_Required",
          "R3_Understanding_or_Explanation_Required",
          "R4_Validation_or_Evidence_Required"],
    "A": ["A1_Not_Mentioned", "A2_General_IP_or_Source_Responsibility",
          "A3_License_or_Source_Compatibility_Required",
          "A4_Documented_Rights_or_Provenance_Required"],
    "C": ["C1_No_Substantive_Restriction", "C2_Form_Restricted",
          "C3_Context_or_Channel_Restricted", "C4_Generally_Excluded"],
    "E": ["E1_Not_Mentioned", "E2_Remediation_or_Review_Limitation",
          "E3_Current_Submission_Rejected_or_Closed",
          "E4_Future_Participation_Restricted"],
}
CONFIDENCE = ["low", "medium", "high"]
SCOPE = ["in_scope", "partly_in_scope", "out_of_scope"]
SEP = "=" * 72


def _body(text: str) -> str:
    return re.sub(r"\s+", " ", text.split(SEP, 1)[-1]).strip().lower()


def dedupe(texts: dict[str, str], threshold: float) -> tuple[list[str], list[dict]]:
    ids = sorted(texts)
    body = {i: _body(texts[i]) for i in ids}
    dropped: dict[str, str] = {}
    log = []
    for a, b in itertools.combinations(ids, 2):
        if a in dropped or b in dropped:
            continue
        sm = difflib.SequenceMatcher(None, body[a], body[b])
        if sm.quick_ratio() < threshold:
            continue
        r = sm.ratio()
        if r >= threshold:
            dropped[b] = a
            log.append({"kept": a, "dropped": b, "similarity": round(r, 3)})
    return [i for i in ids if i not in dropped], log


def published_levels(irr: Path) -> dict[str, set[tuple[str, int]]]:
    out: dict[str, set] = {}
    for f in ("coder_a_labels.csv", "coder_b_labels.csv"):
        df = pd.read_csv(irr / f)
        for _, r in df.iterrows():
            for d, long in zip(DIMS, ["transparency", "responsibility",
                                      "attribution", "constraints", "enforcement"]):
                col = d if d in df.columns else long
                m = re.search(r"[1-4]", str(r[col]))
                if m:
                    out.setdefault(r.sample_id, set()).add((d, int(m.group())))
    return out


def select(pool: list[str], has: dict[str, set], n: int, k: int,
           rng: random.Random) -> tuple[list[str], list[dict]]:
    pairs = sorted({p for i in pool for p in has.get(i, set())},
                   key=lambda p: (sum(p in has.get(i, set()) for i in pool), p))
    chosen: list[str] = []
    log = []
    for p in pairs:
        holders = [i for i in pool if p in has.get(i, set())]
        need = min(k, len(holders)) - sum(i in chosen for i in holders)
        cand = [i for i in holders if i not in chosen]
        rng.shuffle(cand)
        for i in cand[:max(0, need)]:
            if len(chosen) < n:
                chosen.append(i)
                log.append({"sample_id": i, "reason": f"coverage {p[0]}{p[1]}"})
    rest = [i for i in pool if i not in chosen]
    rng.shuffle(rest)
    for i in rest[:max(0, n - len(chosen))]:
        chosen.append(i)
        log.append({"sample_id": i, "reason": "random fill"})
    return sorted(chosen), log


def write_workbook(path: Path, rows: list[dict], columns: list[str],
                   lists: dict[str, list[str]], title: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "coding"
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c, "") for c in columns])
    head = PatternFill("solid", fgColor="DDE6F0")
    fill = PatternFill("solid", fgColor="FFF7D6")
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = head
    widths = {"sample_id": 10, "policy_text": 70, "notes": 40, "confidence": 12,
              "scope": 16}
    for j, col in enumerate(columns, 1):
        L = get_column_letter(j)
        ws.column_dimensions[L].width = widths.get(
            col, 34 if col.endswith("_evidence") else 30)
        for i in range(2, len(rows) + 2):
            cell = ws[f"{L}{i}"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col not in ("sample_id", "policy_text"):
                cell.fill = fill
    for i in range(2, len(rows) + 2):
        ws.row_dimensions[i].height = 160
    for col, values in lists.items():
        formula = '"' + ",".join(values) + '"'
        assert len(formula) <= 255, col
        dv = DataValidation(type="list", formula1=formula,
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="Not a codebook label",
                            error="Pick a value from the list.")
        ws.add_data_validation(dv)
        cl = get_column_letter(columns.index(col) + 1)
        dv.add(f"{cl}2:{cl}{len(rows) + 1}")
    ws.freeze_panes = "C2"
    wb.properties.title = title
    wb.save(path)


def emit(out: Path, rnd: str, rows: list[dict], columns: list[str],
         lists: dict[str, list[str]], coders: list[str]) -> None:
    pd.DataFrame(rows)[columns].to_csv(out / f"{rnd}_sheet.csv", index=False)
    cdir = out / "coders"
    tdir = cdir / f"texts_{rnd}"
    tdir.mkdir(parents=True, exist_ok=True)
    for r in rows:
        (tdir / f"{r['sample_id']}.txt").write_text(r["policy_text"], encoding="utf-8")
    for name in coders:
        write_workbook(cdir / f"{name}_{rnd}.xlsx", rows, columns, lists,
                       f"{name} {rnd}")
    print(f"{len(rows)} rows -> {out}/{rnd}_sheet.csv; workbooks for "
          f"{', '.join(coders)} in {cdir}/; plain texts in {tdir}/")


def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="mode", required=True)
    c = sub.add_parser("calibration")
    c.add_argument("--package", required=True)
    c.add_argument("--n", type=int, default=20)
    c.add_argument("--min-per-level", type=int, default=3,
                   help="each level must appear in at least this many texts")
    c.add_argument("--dedupe-threshold", type=float, default=0.90)
    c.add_argument("--seed", type=int, default=20261005)
    c2 = sub.add_parser("calibration2",
                        help="second calibration round from the reserve, only "
                             "if a dimension stayed below alpha 0.667")
    c2.add_argument("--package", required=True)
    c2.add_argument("--n", type=int, default=10)
    c2.add_argument("--seed", type=int, default=20261012)
    s = sub.add_parser("subject")
    s.add_argument("--units", required=True, help="policy_units/ directory from s2")
    s.add_argument("--redact", nargs="*", default=[])
    s.add_argument("--combine", nargs="*", default=[],
                   help="unit files to merge into one policy unit, e.g. the two "
                        "files that entered in the same commit")
    for p in (c, c2, s):
        p.add_argument("--coders", nargs="+", default=["coder1", "coder2", "coder3"])
        p.add_argument("--out", default="out/coding")
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    label_cols = {d: LEVELS[d] for d in DIMS}

    if a.mode == "calibration2":
        irr = Path(a.package) / "data/policy/irr"
        res_file = out / "calibration_reserve.txt"
        if not res_file.exists():
            raise SystemExit(f"{res_file} not found — run the calibration mode first")
        reserve = [x.strip() for x in res_file.read_text().split() if x.strip()]
        pick = sorted(random.Random(a.seed).sample(reserve, min(a.n, len(reserve))))
        rows = [{"sample_id": i,
                 "policy_text": (irr / "sample_texts" / f"{i}.txt").read_text(
                     encoding="utf-8", errors="replace"),
                 **{d: "" for d in DIMS}, "confidence": "", "notes": ""}
                for i in pick]
        cols = ["sample_id", "policy_text", *DIMS, "confidence", "notes"]
        emit(out, "calibration2", rows, cols,
             {**label_cols, "confidence": CONFIDENCE}, a.coders)
        print(f"round-2 calibration: {len(pick)} reserve texts (seed {a.seed}): "
              + " ".join(pick))
        return 0

    if a.mode == "calibration":
        irr = Path(a.package) / "data/policy/irr"
        texts = {p.stem: p.read_text(encoding="utf-8", errors="replace")
                 for p in (irr / "sample_texts").glob("R*.txt")}
        pool, dup_log = dedupe(texts, a.dedupe_threshold)
        has = published_levels(irr)
        pick, sel_log = select(pool, has, a.n, a.min_per_level,
                               random.Random(a.seed))
        reserve = [i for i in pool if i not in pick]
        rows = [{"sample_id": i, "policy_text": texts[i], **{d: "" for d in DIMS},
                 "confidence": "", "notes": ""} for i in pick]
        cols = ["sample_id", "policy_text", *DIMS, "confidence", "notes"]
        emit(out, "calibration", rows, cols,
             {**label_cols, "confidence": CONFIDENCE}, a.coders)
        cover = {f"{d}{l}": sum((d, l) in has.get(i, set()) for i in pick)
                 for d in DIMS for l in range(1, 5)}
        avail = {f"{d}{l}": sum((d, l) in has.get(i, set()) for i in pool)
                 for d in DIMS for l in range(1, 5)}
        words = {i: len(texts[i].split(SEP, 1)[-1].split()) for i in pick}
        log = {"seed": a.seed, "n": len(pick), "min_per_level": a.min_per_level,
               "pool_after_dedupe": len(pool), "texts_total": len(texts),
               "duplicates_dropped": dup_log, "selected": pick,
               "selection_steps": sel_log, "reserve": reserve,
               "texts_per_level_selected": cover,
               "texts_per_level_available": avail,
               "words_selected_total": sum(words.values()),
               "words_per_text": words}
        (out / "calibration_selection_PRIVATE.json").write_text(json.dumps(log, indent=2))
        (out / "calibration_reserve.txt").write_text("\n".join(reserve) + "\n")
        print(f"dropped {len(dup_log)} duplicate(s): "
              + ", ".join(f"{x['dropped']}={x['kept']}" for x in dup_log))
        print(f"selected {len(pick)} of {len(pool)} unique texts (seed {a.seed}):")
        print("  " + " ".join(pick))
        print(f"reading load: {sum(words.values())} words; "
              f"reserve of {len(reserve)} texts -> {out}/calibration_reserve.txt")
        short = [k for k in cover if cover[k] < min(a.min_per_level, avail[k])]
        if short:
            print("  ! levels below target (n too small):", ", ".join(short))
        return 0

    files = sorted(Path(a.units).glob("*.txt"))
    texts = {f.name: f.read_text(encoding="utf-8") for f in files}
    units = []
    if a.combine:
        missing = [n for n in a.combine if n not in texts]
        if missing:
            raise SystemExit(f"--combine files not found: {missing}; have {list(texts)}")
        units.append(("U1", " + ".join(a.combine),
                      "\n\n".join(texts[n] for n in a.combine)))
        for n in a.combine:
            units.append((f"U{len(units)+1}", n + " (alone)", texts[n]))
    for n in texts:
        if n not in a.combine:
            units.append((f"U{len(units)+1}", n, texts[n]))
    rows, key = [], []
    for uid, src, t in units:
        for word in a.redact:
            t = re.sub(re.escape(word), "[PROJECT]", t, flags=re.I)
        r = {"sample_id": uid, "policy_text": t, "scope": ""}
        for d in DIMS:
            r[d] = ""
            r[f"{d}_evidence"] = ""
        r.update({"confidence": "", "notes": ""})
        rows.append(r)
        key.append({"sample_id": uid, "source": src})
    cols = ["sample_id", "policy_text", "scope"]
    for d in DIMS:
        cols += [d, f"{d}_evidence"]
    cols += ["confidence", "notes"]
    emit(out, "subject", rows, cols,
         {"scope": SCOPE, **label_cols, "confidence": CONFIDENCE}, a.coders)
    pd.DataFrame(key).to_csv(out / "subject_sheet_key_PRIVATE.csv", index=False)
    for k in key:
        print(f"  {k['sample_id']}: {k['source']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
